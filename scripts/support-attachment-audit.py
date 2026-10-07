"""Inventory support attachments and extract bounded text from safe file formats.

The output is private provenance for review. A skipped attachment is never
treated as understood, and extracted text is never an approved resolution.
"""

from __future__ import annotations

import argparse
import csv
import json
import mimetypes
import sqlite3
import zipfile
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree


MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_EXTRACTED_CHARS = 12_000
SNIPPET_CHARS = 300
TEXT_EXTENSIONS = {".txt", ".text", ".md", ".csv", ".json", ".log"}
MEDIA_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".heic",
    ".mp4", ".mov", ".avi", ".webm", ".mkv", ".mp3", ".ogg", ".wav",
    ".opus", ".tgs", ".webp",
}


class VisibleHtml(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style"} and self.hidden:
            self.hidden -= 1

    def handle_data(self, data):
        if not self.hidden and data.strip():
            self.parts.append(data.strip())


def office_text(path: Path, extension: str) -> str:
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(entry.file_size for entry in entries) > 32 * 1024 * 1024:
            raise ValueError("office archive exceeds extraction limit")
        if extension == ".docx":
            names = ["word/document.xml"]
        else:
            names = sorted(name for name in archive.namelist()
                           if name.startswith("ppt/slides/slide") and name.endswith(".xml"))[:100]
        parts = []
        for name in names:
            if name not in archive.namelist():
                continue
            root = ElementTree.fromstring(archive.read(name))
            parts.extend(element.text for element in root.iter()
                         if element.tag.endswith("}t") and element.text)
            if sum(map(len, parts)) >= MAX_EXTRACTED_CHARS:
                break
        return "\n".join(parts)[:MAX_EXTRACTED_CHARS]


def har_summary(path: Path) -> str:
    """Keep request path and status, excluding headers, cookies, bodies and query strings."""
    archive = json.loads(path.read_text(encoding="utf-8-sig"))
    entries = archive.get("log", {}).get("entries", [])
    if not isinstance(entries, list):
        raise ValueError("invalid HAR entries")
    lines = []
    for entry in entries[:500]:
        request = entry.get("request", {})
        response = entry.get("response", {})
        url = urlsplit(request.get("url", ""))
        route = f"{url.scheme}://{url.hostname or ''}{url.path}" if url.scheme else ""
        lines.append(f"{request.get('method', '')} {route} -> {response.get('status', '')}")
        if sum(map(len, lines)) >= MAX_EXTRACTED_CHARS:
            break
    return "\n".join(lines)[:MAX_EXTRACTED_CHARS]


def zip_summary(path: Path) -> str:
    """List archive members without unpacking or executing them."""
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(entry.file_size for entry in entries) > 32 * 1024 * 1024:
            raise ValueError("zip archive exceeds inventory limit")
        return "\n".join(f"{entry.filename} ({entry.file_size} bytes)"
                         for entry in entries[:200])[:MAX_EXTRACTED_CHARS]


def safe_attachment_path(export_root: Path, href: str) -> Path | None:
    """Resolve only local paths contained in the export; reject encoded escapes."""
    try:
        parsed = urlsplit(href)
    except ValueError:
        return None
    if parsed.scheme or parsed.netloc or not parsed.path:
        return None
    decoded = unquote(parsed.path)
    if "\\" in decoded or "\x00" in decoded or decoded.startswith("/"):
        return None
    parts = decoded.split("/")
    if not parts or any(part in {"..", ".", ""} for part in parts):
        return None
    root = export_root.resolve()
    candidate = root.joinpath(*parts).resolve()
    if not candidate.is_relative_to(root):
        return None
    return candidate


def extract_text(path: Path) -> tuple[str, str]:
    """Return status and bounded text. Never execute or unpack the input."""
    extension = path.suffix.lower()
    if extension in MEDIA_EXTENSIONS:
        return "unprocessed_media", ""
    if extension not in TEXT_EXTENSIONS | {".pdf", ".xlsx", ".docx", ".pptx", ".html", ".htm", ".har", ".zip"}:
        return "unsupported_type", ""
    if not path.is_file():
        return "missing_file", ""
    try:
        if path.stat().st_size > MAX_FILE_BYTES:
            return "over_size_limit", ""
        if extension in TEXT_EXTENSIONS:
            text = path.read_text(encoding="utf-8-sig", errors="replace")
        elif extension in {".html", ".htm"}:
            parser = VisibleHtml()
            parser.feed(path.read_text(encoding="utf-8-sig", errors="replace"))
            text = "\n".join(parser.parts)
        elif extension in {".docx", ".pptx"}:
            text = office_text(path, extension)
        elif extension == ".har":
            text = har_summary(path)
        elif extension == ".zip":
            text = zip_summary(path)
        elif extension == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                return "parser_unavailable", ""
            reader = PdfReader(str(path), strict=True)
            text = "\n".join(page.extract_text() or "" for page in reader.pages[:100])
        else:
            try:
                from openpyxl import load_workbook
            except ImportError:
                return "parser_unavailable", ""
            with zipfile.ZipFile(path) as archive:
                if len(archive.infolist()) > 1000 or sum(
                    entry.file_size for entry in archive.infolist()
                ) > 32 * 1024 * 1024:
                    return "over_size_limit", ""
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                lines = []
                for sheet in workbook.worksheets[:20]:
                    lines.append(f"Sheet: {sheet.title}")
                    for row in sheet.iter_rows(max_row=1000, max_col=40, values_only=True):
                        lines.append("\t".join("" if value is None else str(value) for value in row))
                        if sum(map(len, lines)) >= MAX_EXTRACTED_CHARS:
                            break
            finally:
                workbook.close()
            text = "\n".join(lines)
        return ("extracted" if text.strip() else "no_extractable_text"), text[:MAX_EXTRACTED_CHARS]
    except Exception:
        # Third-party parsers may raise format-specific exceptions. Record only
        # the status; exception text can contain private paths or document text.
        return "unreadable", ""


def inventory(connection: sqlite3.Connection, export_root: Path) -> dict:
    connection.row_factory = sqlite3.Row
    rows = connection.execute(
        "SELECT m.revision_id,m.source_message_id,f.file_name,m.media_refs_json "
        "FROM message_revisions m JOIN source_files f ON f.source_file_id=m.source_file_id "
        "WHERE m.corpus_id='support' AND f.source_file_id=("
        " SELECT MAX(latest.source_file_id) FROM source_files latest "
        " WHERE latest.corpus_id=f.corpus_id AND latest.file_name=f.file_name) "
        "ORDER BY m.revision_id"
    ).fetchall()
    attachments = []
    message_ids_with_attachments = set()
    for row in rows:
        refs = json.loads(row["media_refs_json"])
        if not isinstance(refs, list) or not all(isinstance(ref, str) for ref in refs):
            raise ValueError(f"invalid media refs on revision {row['revision_id']}")
        for href in refs:
            message_ids_with_attachments.add(row["revision_id"])
            path = safe_attachment_path(export_root, href)
            extension = PurePosixPath(urlsplit(href).path).suffix.lower()
            if path is None:
                status, extracted = "unsafe_path", ""
                relative_path = None
            else:
                status, extracted = extract_text(path)
                relative_path = path.relative_to(export_root.resolve()).as_posix()
            attachments.append({
                "revision_id": row["revision_id"],
                "source_message_id": row["source_message_id"],
                "source_file": row["file_name"],
                "href": href,
                "relative_path": relative_path,
                "extension": extension or "(none)",
                "mime_guess": mimetypes.guess_type(href)[0],
                "status": status,
                "extracted_text": extracted,
                "text_snippet": " ".join(extracted.split())[:SNIPPET_CHARS],
            })
    if len(message_ids_with_attachments) != len({a["revision_id"] for a in attachments}):
        raise ValueError("attachment coverage mismatch")
    return {
        "report_kind": "private_support_attachment_inventory",
        "disclaimer": "Only extracted text is represented; media and unreadable files remain unprocessed. Extraction does not verify a claim or approve a fix.",
        "current_support_message_count": len(rows),
        "messages_with_attachments_count": len(message_ids_with_attachments),
        "attachment_reference_count": len(attachments),
        "status_counts": dict(sorted(Counter(a["status"] for a in attachments).items())),
        "extension_counts": dict(sorted(Counter(a["extension"] for a in attachments).items())),
        "attachments": attachments,
    }


def write_report(report: dict, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "support-attachments.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    fields = [
        "revision_id", "source_message_id", "source_file", "href", "relative_path",
        "extension", "mime_guess", "status", "text_snippet",
    ]
    with (output_dir / "support-attachments.tsv").open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fields, delimiter="\t")
        writer.writeheader()
        for attachment in report["attachments"]:
            writer.writerow({key: attachment[key] for key in fields})


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("database", type=Path)
    parser.add_argument("export_root", type=Path)
    parser.add_argument(
        "--output-dir", type=Path,
        default=Path(".raydium-debugger/attachment-audit"),
    )
    args = parser.parse_args()
    database = args.database.resolve()
    export_root = args.export_root.resolve()
    if not database.is_file() or not export_root.is_dir():
        parser.error("database and export_root must exist")
    with sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True) as connection:
        report = inventory(connection, export_root)
    write_report(report, args.output_dir)
    print(
        f"Inventoried {report['attachment_reference_count']} attachment references "
        f"in {report['messages_with_attachments_count']} of "
        f"{report['current_support_message_count']} current support messages. "
        f"Private report: {args.output_dir / 'support-attachments.json'}"
    )


if __name__ == "__main__":
    main()
